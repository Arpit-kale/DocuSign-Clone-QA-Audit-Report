#!/usr/bin/env python3
"""Build bug_tracker_<Name>.xlsx and summary_note.pdf from bugs.json + summary_note.md.

Usage:
  python scripts/build_deliverables.py --name "QA"
  python scripts/build_deliverables.py --name "QA" --evidence-base https://github.com/<user>/<repo>/blob/main/
  python scripts/build_deliverables.py --template          # blank template with one example row
"""
import argparse, json, os, re, sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COLS = [  # (header, json key, width)
    ("Bug ID", "id", 10), ("Title", "title", 42), ("Category", "category", 16),
    ("Severity", "severity", 9), ("Module / Screen", "module", 22),
    ("Steps to Reproduce", "steps", 48), ("Expected Behavior", "expected", 36),
    ("Actual Behavior", "actual", 36), ("Evidence", "evidence", 30),
    ("Browser + OS", "browser_os", 22), ("Viewport", "viewport", 11),
    ("Reproducibility", "repro", 14), ("Date Found", "date", 12),
    ("Notes / Hypothesis", "notes", 36),
]
SEVERITIES = ["P0", "P1", "P2", "P3", "P4"]
SEV_NAMES = {"P0": "Blocker", "P1": "Critical", "P2": "Major", "P3": "Minor", "P4": "Trivial"}
SEV_FILL = {"P0": "C00000", "P1": "FF6B6B", "P2": "FFC000", "P3": "9BC2E6", "P4": "D9D9D9"}
CATEGORIES = ["UI / Pixel", "Functional", "Validation", "Navigation", "Data / Seeded",
              "Edge case", "Error handling", "Content / Copy"]
EXAMPLE = {
    "id": "BUG-001", "title": "Send flow: envelope can be sent with zero recipients, no validation shown",
    "category": "Validation", "severity": "P1", "module": "Prepare → Send",
    "steps": "1. Open http://13.207.185.159/ and sign in ([USERNAME] / [PASSWORD]).\n2. Dashboard → Start → Upload sample.pdf.\n3. Leave Recipients empty.\n4. Click Next.\n5. Click Send.",
    "expected": "In DocuSign, Send is blocked and an inline error asks for at least one recipient.",
    "actual": "Envelope is created and appears in Manage with status \"Sent\" and no recipient.",
    "evidence": "evidence/BUG-001_send_zero-recipients.mp4", "repro": "5/5",
    "notes": "EXAMPLE ROW — format reference only, delete before submitting. Validation likely runs on the Recipients screen only, not on submit.",
}

FONT = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")


def load_json(path, default):
    if not os.path.exists(path):
        return default
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def sort_key(b):
    return (b.get("module", ""), SEVERITIES.index(b["severity"]) if b.get("severity") in SEVERITIES else 9, b.get("id", ""))


def validate(bugs):
    errs, seen = [], set()
    for i, b in enumerate(bugs):
        tag = b.get("id", f"row {i+1}")
        for k in ("id", "title", "category", "severity", "module", "steps", "expected", "actual", "evidence"):
            if not str(b.get(k, "")).strip():
                errs.append(f"{tag}: missing '{k}'")
        if b.get("severity") not in SEVERITIES:
            errs.append(f"{tag}: severity must be one of {SEVERITIES}")
        if b.get("id") in seen:
            errs.append(f"{tag}: duplicate id")
        seen.add(b.get("id"))
        ev = b.get("evidence", "")
        if ev and not ev.startswith("http") and not os.path.exists(os.path.join(ROOT, ev)):
            errs.append(f"{tag}: evidence file not found: {ev}")
    return errs


def build_xlsx(bugs, cfg, out, evidence_base):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bug Tracker"
    for c, (h, _, w) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"

    for r, b in enumerate(bugs, 2):
        row = dict(b)
        row.setdefault("browser_os", cfg.get("browser_os", ""))
        row.setdefault("viewport", cfg.get("viewport", "1280x960"))
        row.setdefault("date", cfg.get("date", date.today().strftime("%d %b %Y")))
        for c, (_, key, _) in enumerate(COLS, 1):
            val = row.get(key, "")
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if key == "evidence" and val:
                url = val if val.startswith("http") or not evidence_base else evidence_base.rstrip("/") + "/" + val.lstrip("./")
                cell.hyperlink = url
                cell.value = url if evidence_base else val
                cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")
            if key == "severity" and val in SEV_FILL:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF" if val in ("P0", "P1") else "000000")

    last = max(len(bugs) + 1, 2)
    end = 500  # leave room to add rows by hand
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLS)).column_letter}{last}"

    sev_col = "D"
    for s, color in SEV_FILL.items():
        ws.conditional_formatting.add(f"{sev_col}2:{sev_col}{end}",
                                      FormulaRule(formula=[f'${sev_col}2="{s}"'], fill=PatternFill("solid", fgColor=color)))
    dv_sev = DataValidation(type="list", formula1='"' + ",".join(SEVERITIES) + '"', allow_blank=True)
    dv_cat = DataValidation(type="list", formula1="=Lists!$A$2:$A$" + str(len(CATEGORIES) + 1), allow_blank=True)
    ws.add_data_validation(dv_sev); dv_sev.add(f"D2:D{end}")
    ws.add_data_validation(dv_cat); dv_cat.add(f"C2:C{end}")

    # Summary sheet (live formulas)
    sm = wb.create_sheet("Summary")
    sm["A1"] = "Bug count by severity"; sm["A1"].font = Font(name=FONT, bold=True, size=12)
    sm.append(["Severity", "Name", "Count"])
    for s in SEVERITIES:
        sm.append([s, SEV_NAMES[s], f"=COUNTIF('Bug Tracker'!$D$2:$D${end},A{sm.max_row + 1})"])
    tot = sm.max_row + 1
    sm.append(["Total", "", f"=SUM(C3:C{tot - 1})"])
    sm.append([])
    sm.append(["Bug count by category"])
    sm.cell(row=sm.max_row, column=1).font = Font(name=FONT, bold=True, size=12)
    sm.append(["Category", "", "Count"])
    cat_head = sm.max_row
    for cat in CATEGORIES:
        sm.append([cat, "", f"=COUNTIF('Bug Tracker'!$C$2:$C${end},A{sm.max_row + 1})"])
    bold_rows = {2, tot, cat_head}
    titles = {1, cat_head - 1}
    for row in sm.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=12 if cell.row in titles else 10,
                             bold=cell.row in bold_rows or cell.row in titles)
    sm.column_dimensions["A"].width = 18; sm.column_dimensions["B"].width = 12; sm.column_dimensions["C"].width = 10

    # Lists + legend
    ls = wb.create_sheet("Lists")
    ls["A1"] = "Categories"; ls["C1"] = "Severity"; ls["D1"] = "Name"; ls["E1"] = "Meaning"
    for i, cat in enumerate(CATEGORIES, 2):
        ls.cell(row=i, column=1, value=cat)
    meaning = {
        "P0": "Stops the user cold: crash, data loss, login broken, major flow unusable.",
        "P1": "Major feature broken for many users but a workaround exists, or severe UI breakage on a core screen.",
        "P2": "Functional bug on a common flow that is ugly/confusing but does not block, or pixel issue on a primary screen.",
        "P3": "Edge-case bug, small visual issue, uncommon path. Low user impact.",
        "P4": "Cosmetic: typo, spacing off by a pixel, easter egg.",
    }
    for i, s in enumerate(SEVERITIES, 2):
        ls.cell(row=i, column=3, value=s); ls.cell(row=i, column=4, value=SEV_NAMES[s]); ls.cell(row=i, column=5, value=meaning[s])
    ls.cell(row=9, column=1, value="How to use: one bug per row on 'Bug Tracker'. Pick Category and Severity from the dropdowns. "
                                  "Browser + OS, Viewport and Date default from config.json. Summary counts update automatically.")
    for row in ls.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=cell.row == 1)
    ls.column_dimensions["A"].width = 18; ls.column_dimensions["E"].width = 90

    wb.save(out)


def md_to_pdf(md_path, out, title):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem

    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.5)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=14, leading=17, spaceAfter=4, textColor="#1F3864")
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=6, spaceAfter=2, textColor="#1F3864")

    def inline(t):
        t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
        t = re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", t)
        return t

    flow, bullets, numbered = [], [], False

    def flush():
        nonlocal bullets
        if bullets:
            flow.append(ListFlowable([ListItem(Paragraph(inline(b), body), leftIndent=10) for b in bullets],
                                     bulletType="1" if numbered else "bullet", leftIndent=12, bulletFontSize=8))
            bullets = []

    with open(md_path, encoding="utf-8") as f:
        for line in f:
            s = line.rstrip()
            if s.startswith("<!--") or not s.strip():
                flush(); continue
            m_num = re.match(r"^\d+\.\s+(.*)", s)
            if s.startswith("# "):
                flush(); flow.append(Paragraph(inline(s[2:]), h1))
            elif s.startswith("## "):
                flush(); flow.append(Paragraph(inline(s[3:]), h2))
            elif s.lstrip().startswith(("- ", "* ")):
                if numbered: flush()
                numbered = False; bullets.append(s.lstrip()[2:])
            elif m_num:
                if not numbered: flush()
                numbered = True; bullets.append(m_num.group(1))
            else:
                flush(); flow.append(Paragraph(inline(s), body))
    flush()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=title)
    doc.build(flow)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="YourName")
    ap.add_argument("--evidence-base", default="")
    ap.add_argument("--template", action="store_true", help="write a blank tracker with one example row")
    ap.add_argument("--skip-checks", action="store_true")
    a = ap.parse_args()

    cfg = load_json(os.path.join(ROOT, "config.json"), {})
    slug = re.sub(r"\s+", "", a.name)

    if a.template:
        out = os.path.join(ROOT, "bug_tracker_TEMPLATE.xlsx")
        build_xlsx([EXAMPLE], cfg, out, "")
        print("wrote", out); return

    bugs = load_json(os.path.join(ROOT, "bugs.json"), [])
    if not bugs:
        sys.exit("bugs.json is empty — log bugs first (see bugs.example.json).")
    errs = [] if a.skip_checks else validate(bugs)
    if errs:
        print("Fix these before building:\n  " + "\n  ".join(errs)); sys.exit(1)
    bugs.sort(key=sort_key)

    xlsx = os.path.join(ROOT, f"bug_tracker_{slug}.xlsx")
    build_xlsx(bugs, cfg, xlsx, a.evidence_base)
    print("wrote", xlsx)

    md = os.path.join(ROOT, "summary_note.md")
    if os.path.exists(md):
        pdf = os.path.join(ROOT, "summary_note.pdf")
        md_to_pdf(md, pdf, f"QA Summary — {a.name}")
        print("wrote", pdf)

    counts = {s: sum(1 for b in bugs if b["severity"] == s) for s in SEVERITIES}
    print("severity spread:", counts, "total:", len(bugs))


if __name__ == "__main__":
    main()
